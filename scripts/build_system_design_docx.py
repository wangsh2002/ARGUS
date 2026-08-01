from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "ARGUS_SYSTEM_DESIGN.md"
OUTPUT = ROOT / "docs" / "ARGUS_SYSTEM_DESIGN.docx"

FONT_CN = "PingFang SC"
FONT_LATIN = "Arial"
RED = RGBColor(198, 43, 55)
DARK = RGBColor(36, 43, 52)
MUTED = RGBColor(101, 111, 122)
LIGHT = "F4F6F8"
LIGHT_RED = "FBEDEF"
BORDER = "D9DEE5"
WHITE = RGBColor(255, 255, 255)
USABLE_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(
    run,
    *,
    size: float = 11,
    color: RGBColor = DARK,
    bold: bool | None = None,
    italic: bool | None = None,
    name: str = FONT_CN,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size=4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def set_paragraph_border(paragraph, *, color="C62B37", size=14) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "7")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.16

    heading_tokens = {
        "Heading 1": (17, RED, 16, 8),
        "Heading 2": (13.5, DARK, 12, 6),
        "Heading 3": (11.5, RGBColor(68, 76, 87), 9, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.16


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_inline_markdown(paragraph, text: str, *, base_size=10.5, base_color=DARK) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=base_size, color=base_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, color=base_color, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(
                run,
                size=max(9, base_size - 0.5),
                color=RGBColor(118, 40, 49),
                name="Menlo",
            )
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=base_size, color=base_color)


def add_body_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.16
    add_inline_markdown(paragraph, text)


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [USABLE_DXA])
    set_table_borders(table, color="E6C5C9", size=4)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_RED)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.16
    add_inline_markdown(paragraph, text, base_size=10.5, base_color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        rows.append([cell.strip() for cell in line.strip().strip("|").split("|")])
    return rows


def is_separator_row(row: list[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in row)


def table_widths(rows: list[list[str]]) -> list[int]:
    columns = len(rows[0])
    lengths = []
    for column in range(columns):
        maximum = max(len(row[column]) if column < len(row) else 0 for row in rows)
        lengths.append(max(8, min(maximum, 42)))
    if columns == 2:
        first_ratio = min(0.34, max(0.22, lengths[0] / sum(lengths)))
        first = int(USABLE_DXA * first_ratio)
        return [first, USABLE_DXA - first]
    total = sum(lengths)
    widths = [int(USABLE_DXA * value / total) for value in lengths]
    widths[-1] += USABLE_DXA - sum(widths)
    return widths


def add_table(doc: Document, raw_lines: list[str]) -> None:
    rows = parse_table(raw_lines)
    if len(rows) > 1 and is_separator_row(rows[1]):
        rows.pop(1)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = None
    set_table_geometry(table, table_widths(rows))
    set_table_borders(table)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            set_cell_shading(cell, "C62B37" if row_index == 0 else ("F8F9FA" if row_index % 2 == 0 else "FFFFFF"))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.08
            add_inline_markdown(
                paragraph,
                value,
                base_size=9.2 if len(rows[0]) >= 3 else 9.6,
                base_color=WHITE if row_index == 0 else DARK,
            )
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [USABLE_DXA])
    set_table_borders(table, color="E1E5EA", size=4)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.05
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        if index < len(lines) - 1:
            run.add_break()
        set_run_font(run, size=8.6, color=RGBColor(55, 63, 73), name="Menlo")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("APT 技战术数据服务平台")
    set_run_font(run, size=12, color=RED, bold=True)
    kicker.paragraph_format.space_after = Pt(16)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ARGUS")
    set_run_font(run, size=32, color=DARK, bold=True)
    title.paragraph_format.space_after = Pt(5)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("前端、后端与数据库总体设计说明")
    set_run_font(run, size=18, color=DARK, bold=True)
    subtitle.paragraph_format.space_after = Pt(28)

    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = rule.add_run("从真实来源到可解释攻击路径与行动建议")
    set_run_font(run, size=11, color=MUTED)
    set_paragraph_border(rule)
    rule.paragraph_format.space_after = Pt(30)

    metadata = [
        ("文档版本", "1.0"),
        ("编制日期", "2026-07-31"),
        ("适用阶段", "比赛展示版 / 后续工程化设计基线"),
        ("文档状态", "前端以当前实现为准；后端按实现状态分级描述"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    set_table_geometry(table, [2100, USABLE_DXA - 2100])
    set_table_borders(table, color="E4E7EB", size=4)
    for row_index, (label, value) in enumerate(metadata):
        for column_index, text in enumerate((label, value)):
            cell = table.cell(row_index, column_index)
            set_cell_shading(cell, "F4F6F8" if column_index == 0 else "FFFFFF")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(text)
            set_run_font(
                run,
                size=10,
                color=MUTED if column_index == 0 else DARK,
                bold=column_index == 0,
            )

    for _ in range(3):
        doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("内部设计基线 · 所有结论以可追溯证据为前提")
    set_run_font(run, size=9.5, color=MUTED)
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    heading = doc.add_paragraph(style="Heading 1")
    heading.add_run("文档阅读说明")
    add_body_paragraph(
        doc,
        "本文档将平台能力分为“已实现、部分实现、设计确定”三类。前端章节以当前界面和交互为准；后端与数据库章节同时描述现有代码、比赛版边界和后续工程化方案。",
    )
    add_callout(
        doc,
        "核心结论：ARGUS 的价值不是单纯展示图谱，而是把碎片线索转化为可追踪的攻击过程、补全与预测依据、处置建议和可审计证据链。",
    )
    heading = doc.add_paragraph(style="Heading 2")
    heading.add_run("章节导航")
    items = [
        "第 1–5 章：产品定位、用户、原则与总体架构",
        "第 6–7 章：前端页面、交互与状态设计",
        "第 8–10 章：后端流程、算法与标准本体",
        "第 11–13 章：逻辑模型、SQLite 物理表与 API",
        "第 14–17 章：数据来源、演示闭环、安全与质量",
        "第 18–20 章：实现状态、路线和验收标准",
    ]
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        add_inline_markdown(paragraph, item)
    doc.add_page_break()


def add_document_body(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    start = next(index for index, line in enumerate(lines) if line.startswith("## 1. 文档目的"))
    index = start
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_buffer:
            text = " ".join(part.strip() for part in paragraph_buffer)
            add_body_paragraph(doc, text)
            paragraph_buffer.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("---"):
            flush_paragraph()
            index += 1
            continue

        if stripped.startswith("```"):
            flush_paragraph()
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(doc, code_lines)
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, table_lines)
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            flush_paragraph()
            level = min(len(heading_match.group(1)) - 1, 3)
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            add_inline_markdown(
                paragraph,
                heading_match.group(2),
                base_size={1: 17, 2: 13.5, 3: 11.5}[level],
                base_color=RED if level == 1 else (DARK if level == 2 else RGBColor(68, 76, 87)),
            )
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            callout_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                callout_lines.append(lines[index].strip().lstrip(">").strip())
                index += 1
            add_callout(doc, " ".join(callout_lines))
            continue

        bullet_match = re.match(r"^-\s+(.+)$", stripped)
        if bullet_match:
            flush_paragraph()
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, bullet_match.group(1))
            index += 1
            continue

        number_match = re.match(r"^\d+\.\s+(.+)$", stripped)
        if number_match:
            flush_paragraph()
            paragraph = doc.add_paragraph(style="List Number")
            add_inline_markdown(paragraph, number_match.group(1))
            index += 1
            continue

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()


def configure_page(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.82)
        section.bottom_margin = Inches(0.78)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.36)
        section.footer_distance = Inches(0.36)

        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run("ARGUS 系统设计说明  ·  v1.0")
        set_run_font(run, size=8.5, color=MUTED, bold=True)
        set_paragraph_border(paragraph, color="E2E5E9", size=4)

        footer = section.footer
        add_page_field(footer.paragraphs[0])


def build() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    add_cover(doc)
    add_front_matter(doc)
    add_document_body(doc, markdown)

    properties = doc.core_properties
    properties.title = "ARGUS 前端、后端与数据库总体设计说明"
    properties.subject = "APT 技战术数据服务平台系统设计"
    properties.author = "ARGUS 项目组"
    properties.keywords = "ARGUS, APT, 威胁情报, 攻击路径, STIX, ATT&CK, Attack Flow"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
